#!/usr/bin/env python3
"""
pepsico_portafolio.py — pepsico/oportunidad_portafolio.json

*** REQUIERE UN EXPORT MANUAL, NO SE PUEDE AUTOMATIZAR POR API TODAVIA ***
Fuente: GesCom, Ventas / Reportes / Rumbo Pepsico / LATAM (Desde 04-2024) /
6.1 Censo Tienda perfecta — es un reporte armado a medida para el programa
Tienda Perfecta de Pepsico (relevamiento de portafolio por supervisor), no
encontramos equivalente en la API general de GesCom (tools/gescom.py). Puede
existir un endpoint propio (relevamientos/encuestas) que no se llego a probar
por falta de acceso a un entorno con credenciales — ver README.md.

Tambien necesita el maestro de clientes (Clientes Claude.xlsx: direccion,
localidad y telefono, que el censo no siempre trae completos). Ese maestro
en teoria SI se podria reemplazar por la API (gescom.bajar_clientes(), el
mismo endpoint que ya usa tools/gescom.py) — no se hizo aca todavia porque el
cruce de columnas del censo (orden de columnas fijo, ver mas abajo) habria
que revisarlo junto con eso.

Uso:
    python tools/pepsico_portafolio.py "C:\\ruta\\al\\censo.xlsx" "C:\\ruta\\a\\Clientes Claude.xlsx"
"""
import json
import sys
from pathlib import Path

import openpyxl

DIR = Path(__file__).resolve().parent.parent / "pepsico"
DIAS_MAP = {"Lu": "Lunes", "Ma": "Martes", "Mi": "Miercoles", "Ju": "Jueves", "Vi": "Viernes", "Sa": "Sabado"}
VEND_NAMES = {
    "1": "Maria Noelia Madia", "2": "Julian Nuñez", "3": "Ezequiel Hack",
    "4": "Lucas Giayetto", "5": "Nicolas Martinez", "6": "Camila Martinez",
    "7": "Ramiro Rando", "8": "Martinica Rodriguez", "9": "Ludmila Aylen Martinez",
    "10": "Matias Parada", "11": "Melanie Larrama", "12": "Agustin Pogorzelski",
}


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    censo_path, clientes_path = sys.argv[1], sys.argv[2]

    wbc = openpyxl.load_workbook(clientes_path, read_only=True, data_only=True)
    wsc = wbc["Hoja1"]
    cliente_info = {}
    for r in wsc.iter_rows(min_row=2, values_only=True):
        cliente_info[str(r[0])] = {"direccion": r[2], "localidad": r[3], "telefono": r[4]}
    wbc.close()

    wb = openpyxl.load_workbook(censo_path, read_only=True, data_only=True)
    ws = wb["Principal"]
    detalle = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        (cliente, razon, direccion, vendedor_id, tipologia, subcanal, segmento, pct_portafolio,
         cumplimiento_estandar, rackadecuado, rackposicion, planograma, preciosvisibles,
         rotacionpreventiva, puntuacion, tiendaperfecta_field, supok, ultimafechacenso, frecuencia,
         subcanal_tp, fecha_valid_sup, usuario, estandar_sod, puntaje_sod, cumple_tp_field) = r
        if subcanal_tp != "Si":
            continue
        vend = VEND_NAMES.get(str(vendedor_id))
        if not vend:
            continue
        pct = pct_portafolio if isinstance(pct_portafolio, (int, float)) else 0
        if not (60 <= pct <= 79):
            continue
        info = cliente_info.get(str(cliente), {})
        detalle.append({
            "vendedor": vend, "codigo": cliente, "razon": razon,
            "direccion": info.get("direccion") or direccion, "localidad": info.get("localidad"),
            "telefono": info.get("telefono"), "seg": segmento,
            "dia": DIAS_MAP.get(frecuencia), "pct": pct,
        })

    ruta = DIR / "oportunidad_portafolio.json"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(detalle, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Guardado", ruta, ":", len(detalle), "filas")
    return 0


if __name__ == "__main__":
    sys.exit(main())
