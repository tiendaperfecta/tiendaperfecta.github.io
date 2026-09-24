#!/usr/bin/env python3
"""
pepsico_mn.py — pepsico/{mn_seguimiento,mn_resumen_vendedor}.json

*** REQUIERE 3 EXPORTS MANUALES DE POWER BI, NO SE PUEDE AUTOMATIZAR POR LA
API DE GESCOM — ES OTRO SISTEMA ***
Fuente: Power BI "Mi Negocio +" de Pepsico (no es GesCom, es un reporte propio
de Pepsico embebido en Power BI). No hay credenciales de eso en los secrets
del repo (solo estan las de GesCom), y Power BI no tiene un botn de API
publica accesible sin permisos separados de Pepsico/Microsoft. La unica forma
de sacar esto hoy es entrando al reporte con un browser, filtrando por
Fully Digital / Hibrid / Non Digital (los 3 toggles del visual) y usando
"..." -> Exportar datos -> Datos con diseno actual, una vez por cada filtro.

Automatizar esto de verdad requeriria: (a) que Pepsico habilite acceso a la
API de Power BI para este reporte con un service principal, o (b) automatizar
el navegador (login + clicks) desde la Action, que es fragil y probablemente
contra los terminos de uso — no se hizo ninguna de las dos cosas.

Uso (los 3 archivos son los que bajan del boton "Exportar datos" de Power BI,
uno por cada filtro del visual):
    python tools/pepsico_mn.py fully_digital.xlsx hibrid.xlsx non_digital.xlsx "C:\\ruta\\a\\Clientes Claude.xlsx"
"""
import collections
import json
import sys
from pathlib import Path

import openpyxl

DIR = Path(__file__).resolve().parent.parent / "pepsico"
DIAS = ["Lunes", "Martes", "Miercoles", "Jueves", "Viernes", "Sabado"]
UMBRAL_FULLY_DIGITAL = 0.70


def main():
    if len(sys.argv) < 5:
        print(__doc__)
        return 1
    files = sys.argv[1:4]
    clientes_path = sys.argv[4]

    wbc = openpyxl.load_workbook(clientes_path, read_only=True, data_only=True)
    wsc = wbc["Hoja1"]
    cliente_info = {}
    for r in wsc.iter_rows(min_row=2, values_only=True):
        codigo, direccion, localidad, telefono, seg, ruta = r[0], r[2], r[3], r[4], r[17], r[14]
        dia = None
        if ruta:
            for d in DIAS:
                if ruta.strip().startswith(d):
                    dia = d
                    break
        cliente_info[str(codigo)] = {"dia": dia, "direccion": direccion, "localidad": localidad,
                                      "telefono": telefono, "seg": seg}
    wbc.close()

    master, seen_keys = [], set()
    for path in files:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or len(r) < 16 or not r[1]:
                continue
            (distrib, vend, cliente, universo, clientes_mn, registrados, activos,
             sellout, sellout_mn, vs_total, sellout_kg, sellout_u, drop, visitas, visitas_mn,
             fully_digital_tag) = r
            cod = cliente.split(" - ")[0].strip() if cliente else None
            nombre = cliente.split(" - ", 1)[1].strip() if cliente and " - " in cliente else cliente
            key = (vend, cod)
            if key in seen_keys:
                continue
            seen_keys.add(key)
            ratio = vs_total if vs_total is not None else None
            activo_b = bool(activos)
            fully_digital_70 = bool(activo_b and ratio is not None and ratio >= UMBRAL_FULLY_DIGITAL)
            master.append({
                "vendedor": vend, "cod": cod, "nombre": nombre,
                "universo": bool(universo), "registrado": bool(registrados), "activo": activo_b,
                "sellout": sellout or 0, "sellout_mn": sellout_mn or 0, "ratio": ratio,
                "sellout_kg": sellout_kg or 0, "fully_digital_70": fully_digital_70,
            })

    resumen = collections.defaultdict(lambda: {"universo": 0, "registrado": 0, "sellout_kg": 0.0, "fully_digital": 0})
    for c in master:
        d = resumen[c["vendedor"]]
        if c["universo"]:
            d["universo"] += 1
        if c["registrado"]:
            d["registrado"] += 1
        d["sellout_kg"] += c["sellout_kg"]
        if c["fully_digital_70"]:
            d["fully_digital"] += 1

    seguimiento = []
    for c in master:
        if not c["registrado"] or (c["activo"] and c["fully_digital_70"]):
            continue
        info = cliente_info.get(c["cod"], {})
        motivo = "Sin activar" if not c["activo"] else f"Debajo del {int(UMBRAL_FULLY_DIGITAL*100)}% (no fully digital)"
        seguimiento.append({
            "dia": info.get("dia") or "(sin ruta)", "vendedor": c["vendedor"], "cod": c["cod"],
            "nombre": c["nombre"], "direccion": info.get("direccion"), "localidad": info.get("localidad"),
            "telefono": info.get("telefono"), "seg": info.get("seg"), "activo": c["activo"],
            "ratio": round(c["ratio"] * 100, 1) if c["ratio"] is not None else None,
            "sellout": round(c["sellout"]), "sellout_mn": round(c["sellout_mn"]), "motivo": motivo,
        })

    DIR.mkdir(parents=True, exist_ok=True)
    (DIR / "mn_resumen_vendedor.json").write_text(
        json.dumps({v: resumen[v] for v in resumen}, ensure_ascii=False, indent=2), encoding="utf-8")
    (DIR / "mn_seguimiento.json").write_text(
        json.dumps(seguimiento, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Guardado mn_resumen_vendedor.json y mn_seguimiento.json:", len(seguimiento), "clientes a seguimiento")
    return 0


if __name__ == "__main__":
    sys.exit(main())
